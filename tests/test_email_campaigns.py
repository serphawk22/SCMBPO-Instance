"""
Tests for Bulk Excel Upload campaigns (routers/email_campaigns.py, modules/email_campaign_worker.py) and the
removal of SCMBPO's old Excel campaign routes (D6). Ported from serphawk_crm_seo; the fixtures (in-memory
SQLite, fake SMTP sender, test-only Fernet key, no OpenAI key) come from tests/test_followups.py.

    venv\\Scripts\\python -m unittest discover -s tests -t . -v
"""
import csv
import io
import os
import smtplib
import unittest
from datetime import timedelta
from unittest import mock

from sqlmodel import Session, select

from database import (
    ActivityLog, EmailCampaign, EmailCampaignRecipient, FollowUp, Lead, Notification, SentEmail, Tenant, User,
    UserEmailIntegration,
)
from modules import email_campaign_worker as worker, followups_core as core
from modules.email_crypto import encrypt_app_password
from routers import email_campaigns as campaigns_router
from tests._support import ENGINE, main_client
from tests.test_followups import ADMIN, ADMIN_T2, INTERN, NOW, OTHER_SALES, SALES, ApiTestCase, DBTestCase, FakeSender


class CampaignApiTestCase(ApiTestCase):
    extra_routers = (campaigns_router.router,)


class PersonalizeTests(unittest.TestCase):
    def test_ai_personalization_result_is_validated(self):
        from modules import email_ai
        with mock.patch.object(email_ai, "_chat_json", return_value={"subject": "", "body": "Hi Jane, tailored."}):
            p = worker.personalize("Original subject", "Hi Jane", {"Notes": "likes ads"})
        self.assertEqual(p, {"subject": "Original subject", "body": "Hi Jane, tailored."})
        with mock.patch.object(email_ai, "_chat_json", return_value={"subject": "x", "body": " "}):
            with self.assertRaises(email_ai.EmailAIError):
                worker.personalize("s", "b", {})


# ─────────────────────────────────────────────────────────────────────────────
# Upload parsing + templating
# ─────────────────────────────────────────────────────────────────────────────
def xlsx_bytes(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseAndTemplateTests(unittest.TestCase):
    def test_xlsx_with_detected_columns_and_numeric_cells(self):
        data = xlsx_bytes([["First Name", "Last Name", "Email Address", "Company", "Phone"],
                           ["Jane", "Doe", "jane@acme.com", "Acme", 9876543210],
                           [None, None, None, None, None],
                           ["Bob", "", "bob@beta", "Beta", None]])
        parsed = worker.parse_upload("contacts.xlsx", data)
        self.assertEqual(parsed["columns"], ["First Name", "Last Name", "Email Address", "Company", "Phone"])
        self.assertEqual(len(parsed["rows"]), 2)  # blank row dropped
        self.assertEqual(parsed["rows"][0]["Phone"], "9876543210")  # not "9876543210.0"
        self.assertEqual(parsed["detected"]["email"], "Email Address")
        self.assertEqual(parsed["detected"]["first_name"], "First Name")
        self.assertEqual(parsed["detected"]["company"], "Company")
        self.assertEqual(parsed["stats"], {"rows": 2, "sendable": 1, "missing_email": 0, "invalid_email": 1, "duplicates": 0})
        self.assertIn("phone", parsed["placeholders"])

    def test_csv_with_bom_semicolons_and_email_detected_by_content(self):
        text = "﻿Contact;Org;Address\nJane Doe;Acme;jane@acme.com\nJohn;Acme;JANE@acme.com\n"
        parsed = worker.parse_upload("list.csv", text.encode("utf-8"))
        self.assertEqual(parsed["columns"], ["Contact", "Org", "Address"])
        self.assertEqual(parsed["detected"]["email"], "Address")
        self.assertEqual(parsed["detected"]["name"], "Contact")
        self.assertEqual(parsed["stats"]["duplicates"], 1)

    def test_rejects_bad_files(self):
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("old.xls", b"\xd0\xcf\x11\xe0 legacy")
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("x.csv", b"")
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("x.pdf", b"%PDF")
        with mock.patch.dict(os.environ, {"BULK_EMAIL_MAX_ROWS": "2"}):
            with self.assertRaises(worker.UploadError):
                worker.parse_upload("x.csv", b"email\na@x.com\nb@x.com\nc@x.com\n")

    def test_render_placeholders_fallbacks_and_unknowns(self):
        mapping = {"email": "Email", "name": "Full Name", "company": "Company"}
        ctx = worker.row_context({"Full Name": "Jane Doe", "Email": "j@a.com", "Company": "", "City": "Pune"},
                                 mapping, "Ada")
        text, missing = worker.render("Hi {{ first_name }}, {{company|your team}} in {{City}} — {{sender_name}}", ctx)
        self.assertEqual(text, "Hi Jane, your team in Pune — Ada")
        self.assertEqual(missing, [])
        text, missing = worker.render("Hi {{first_name}}, about {{company}}.", worker.row_context({"Email": "x@y.com"}, mapping))
        self.assertEqual(text, "Hi there, about.")
        self.assertEqual(missing, ["company"])
        self.assertEqual(worker.unknown_placeholders(["{{city}} {{nope}}", "{{first_name}}"], ["City"]), ["nope"])


# ─────────────────────────────────────────────────────────────────────────────
# Campaign API + worker
# ─────────────────────────────────────────────────────────────────────────────
ROWS = [
    {"First Name": "Jane", "Email": "jane@acme.com", "Company": "Acme", "City": "Pune"},
    {"First Name": "", "Email": "new@delta.io", "Company": "Delta", "City": "Delhi"},
    {"First Name": "Dup", "Email": "JANE@acme.com", "Company": "Acme", "City": ""},
    {"First Name": "Bad", "Email": "not-an-email", "Company": "X", "City": ""},
    {"First Name": "Empty", "Email": "", "Company": "Y", "City": ""},
]
COLUMNS = ["First Name", "Email", "Company", "City"]
MAPPING = {"email": "Email", "first_name": "First Name", "company": "Company"}


def campaign_payload(**over):
    body = {"columns": COLUMNS, "rows": ROWS, "mapping": MAPPING, "subject": "Growth ideas for {{company}}",
            "body": "Hi {{first_name}},\n\nWe help businesses in {{city|your city}} grow.\n\nBest regards,\nAda",
            "name": "October outreach", "source_filename": "october.xlsx", "create_leads": True, "followup_days": 3}
    body.update(over)
    return body


class CampaignApiTests(CampaignApiTestCase):
    def create(self, user=ADMIN, **over):
        with mock.patch.object(worker, "start_campaign") as start:
            res = self.req("POST", "/email-campaigns", user=user, json=campaign_payload(**over))
        return res, start

    def run_worker(self, cid, sender=None):
        worker.run_campaign(cid, sender=sender or self.sender, delay=0)

    def test_parse_endpoint_accepts_xlsx_upload(self):
        data = xlsx_bytes([["Name", "Email"], ["Jane Doe", "jane@acme.com"]])
        res = self.req("POST", "/email-campaigns/parse", files={"file": ("list.xlsx", data, "application/octet-stream")})
        self.assertEqual(res.status_code, 200, res.text)
        self.assertEqual(res.json()["detected"]["email"], "Email")
        bad = self.req("POST", "/email-campaigns/parse", files={"file": ("list.xls", b"\xd0\xcf", "application/octet-stream")})
        self.assertEqual(bad.status_code, 400)
        self.assertIn(".xlsx", bad.json()["detail"])

    def test_sender_status_names_the_company_account(self):
        res = self.req("GET", "/email-campaigns/sender").json()
        self.assertEqual((res["configured"], res["address"], res["source"]), (True, "ada@agency-mail.test", "user_integration"))
        self.assertNotIn("password", str(res).lower())

    def test_preview_renders_each_row_and_lists_skips(self):
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(count=5))
        self.assertEqual(res.status_code, 200, res.text)
        data = res.json()
        self.assertEqual(data["sendable"], 2)
        first, second = data["previews"]
        self.assertEqual(first["subject"], "Growth ideas for Acme")
        self.assertTrue(first["body"].startswith("Hi Jane,"))
        self.assertIn("in Pune grow", first["body"])
        self.assertTrue(second["body"].startswith("Hi there,"))
        self.assertTrue(first["body"].endswith("Best Regards,\nAda\nSCM BPO"))  # Email Agent signature
        self.assertEqual(sorted(s["reason"].split()[0] for s in data["skipped"]), ["Duplicate", "No", "“not-an-email”"])

    def test_unknown_placeholder_and_missing_email_column_are_rejected(self):
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(subject="Hi {{nickname}}"))
        self.assertEqual(res.status_code, 400)
        self.assertIn("{{nickname}}", res.json()["detail"])
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(mapping={"email": None}))
        self.assertEqual(res.status_code, 400)

    def test_campaign_sends_personalized_emails_from_the_company_account_and_records_everything(self):
        res, start = self.create()
        self.assertEqual(res.status_code, 200, res.text)
        cid = res.json()["campaign"]["id"]
        start.assert_called_once_with(cid)
        self.assertEqual(res.json()["campaign"]["status"], "Queued")
        self.run_worker(cid)

        self.assertEqual([c["to"] for c in self.sender.calls], ["jane@acme.com", "new@delta.io"])
        self.assertTrue(all(c["from"] == "ada@agency-mail.test" for c in self.sender.calls))
        self.assertEqual(self.sender.calls[0]["subject"], "Growth ideas for Acme")
        self.assertTrue(self.sender.calls[0]["body"].startswith("Hi Jane,"))
        self.assertTrue(self.sender.calls[1]["body"].startswith("Hi there,"))
        self.assertEqual(self.sender.calls[0]["body"].lower().count("best regards"), 1)  # template closing replaced by signature

        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.total, c.sent_count, c.failed_count, c.skipped_count), ("Completed", 5, 2, 0, 3))
            recips = {r.row_number: r for r in s.exec(select(EmailCampaignRecipient)).all()}
            self.assertEqual([recips[i].status for i in range(1, 6)], ["Sent", "Sent", "Skipped", "Skipped", "Skipped"])
            self.assertIsNotNone(recips[1].sent_at)
            self.assertEqual(recips[1].lead_id, 1)  # linked to the existing Acme lead, not duplicated
            sent = s.exec(select(SentEmail).order_by(SentEmail.id)).all()
            self.assertEqual([(e.to_email, e.status, e.tenant_id) for e in sent],
                             [("jane@acme.com", "Sent", 1), ("new@delta.io", "Sent", 1)])
            self.assertEqual(self.sender.calls[0]["tracking_id"], sent[0].id)  # open tracking pixel id
            new_lead = s.get(Lead, recips[2].lead_id)
            self.assertEqual((new_lead.email, new_lead.source, new_lead.company_name, new_lead.owner_id),
                             ("new@delta.io", "Email Campaign", "Delta", ADMIN))
            self.assertEqual(s.get(Lead, 1).status, "Contacted")
            fus = s.exec(select(FollowUp).order_by(FollowUp.id)).all()
            self.assertEqual(len(fus), 2)
            self.assertEqual({f.source for f in fus}, {"campaign"})
            self.assertEqual([f.owner_id for f in fus], [SALES, ADMIN])  # Jane's lead owner; the new lead's owner (sender)
            self.assertEqual(core.as_utc(fus[0].due_at).date(), (NOW + timedelta(days=3)).date())
            self.assertTrue(s.exec(select(Notification).where(Notification.user_id == ADMIN)).first())
            self.assertTrue(s.exec(select(ActivityLog).where(ActivityLog.lead_id == 1)).first())

        detail = self.req("GET", f"/email-campaigns/{cid}?status=Sent").json()
        self.assertEqual([r["email"] for r in detail["recipients"]], ["jane@acme.com", "new@delta.io"])
        export = self.req("GET", f"/email-campaigns/{cid}/export")
        self.assertEqual(export.status_code, 200)
        lines = list(csv.reader(io.StringIO(export.text.lstrip("﻿"))))
        self.assertEqual(lines[0][:5], ["Row", "Recipient email", "Name", "Company", "Status"])
        self.assertEqual([l[4] for l in lines[1:]], ["Sent", "Sent", "Skipped", "Skipped", "Skipped"])

    def test_requires_a_configured_company_account(self):
        with mock.patch.dict(os.environ, {"EMAIL_SENDER": "fixed@server.test", "EMAIL_PASSWORD": "pw",
                                          "SMTP_SERVER": "smtp.server.test"}):
            res, start = self.create(user=OTHER_SALES)  # no integration, tenant has no mailbox
        self.assertEqual(res.status_code, 503)
        self.assertIn("Email Integration", res.json()["detail"])
        start.assert_not_called()

    def test_failed_rows_are_recorded_and_can_be_retried(self):
        res, _ = self.create(followup_days=None)
        cid = res.json()["campaign"]["id"]
        self.run_worker(cid, FakeSender(fail_for={"new@delta.io"}))
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.sent_count, c.failed_count), ("Completed", 1, 1))
            failed = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.status == "Failed")).one()
            self.assertIn("refused this recipient", failed.error)
            self.assertEqual(len(s.exec(select(SentEmail)).all()), 1)  # failed send leaves no sent_emails row
            self.assertIsNone(failed.lead_id)  # ...and no lead created just for the bounced address
            self.assertIsNone(s.exec(select(Lead).where(Lead.email == "new@delta.io")).first())
        with mock.patch.object(worker, "start_campaign") as start:
            retry = self.req("POST", f"/email-campaigns/{cid}/retry")
        self.assertEqual(retry.status_code, 200, retry.text)
        self.assertEqual(retry.json()["requeued"], 1)
        start.assert_called_once_with(cid)
        self.run_worker(cid)
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.sent_count, c.failed_count), ("Completed", 2, 0))

    def test_login_failure_stops_the_whole_campaign(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.run_worker(cid, FakeSender(error=smtplib.SMTPAuthenticationError(535, b"bad credentials")))
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual(c.status, "Failed")
            self.assertIn("app password", c.last_error)
            statuses = sorted(r.status for r in s.exec(select(EmailCampaignRecipient)).all())
            self.assertEqual(statuses.count("Failed"), 1)
            self.assertEqual(statuses.count("Pending"), 1)  # never attempted

    def test_cancel_stops_before_the_next_row(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.assertEqual(self.req("POST", f"/email-campaigns/{cid}/cancel").status_code, 200)
        self.run_worker(cid)
        self.assertEqual(self.sender.calls, [])
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}").json()["campaign"]["pending"], 2)

    def test_interrupted_rows_are_failed_not_resent(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        with Session(ENGINE) as s:
            r = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.row_number == 1)).one()
            r.status, r.claimed_at = "Sending", NOW - timedelta(minutes=30)
            s.add(r)
            s.commit()
        with mock.patch.object(worker, "start_campaign", return_value=True) as start:
            self.assertEqual(worker.resume_campaigns(), 1)
        start.assert_called_once()
        self.assertEqual(start.call_args.args[0], cid)
        self.run_worker(cid)
        self.assertEqual([c["to"] for c in self.sender.calls], ["new@delta.io"])
        with Session(ENGINE) as s:
            r = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.row_number == 1)).one()
            self.assertEqual(r.status, "Failed")
            self.assertIn("Interrupted", r.error)

    def test_campaigns_are_scoped_by_role_and_tenant(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}", user=SALES).status_code, 404)
        self.assertEqual(self.req("GET", "/email-campaigns", user=SALES).json()["items"], [])
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}", user=ADMIN_T2).status_code, 404)
        self.assertEqual(self.req("GET", "/email-campaigns", user=INTERN).status_code, 403)
        self.assertEqual(self.req("GET", "/email-campaigns", user=None).status_code, 401)
        self.assertEqual(len(self.req("GET", "/email-campaigns").json()["items"]), 1)

    def test_demo_trial_limit_counts_every_email(self):
        with Session(ENGINE) as s:
            s.add(User(id=9, tenant_id=1, email="demo@agency.test", password="x", name="Demo", role="Demo"))
            s.add(UserEmailIntegration(tenant_id=1, user_id=9, email="demo@agency-mail.test",
                                       encrypted_app_password=encrypt_app_password("pw"), smtp_host="smtp.x", imap_host="imap.x"))
            t = s.get(Tenant, 1)
            t.is_trial, t.limit_emails, t.usage_emails = True, 5, 4
            s.add(t)
            s.commit()
        res, start = self.create(user=9)
        self.assertEqual(res.status_code, 403)
        self.assertEqual(res.json()["detail"]["error"], "LIMIT_REACHED")
        start.assert_not_called()


# ─────────────────────────────────────────────────────────────────────────────
# D6: the old ExcelCampaign routes in main.py are gone; /email-campaigns is the router's
# ─────────────────────────────────────────────────────────────────────────────
class OldExcelRoutesRemovedTests(DBTestCase):
    def setUp(self):
        super().setUp()
        self.seed()
        self.client = main_client(self)

    def test_email_campaigns_is_served_by_the_new_router(self):
        import main
        paths = {(m, r.path) for r in main.app.routes for m in getattr(r, "methods", ()) or ()}
        self.assertNotIn(("POST", "/email-campaigns/excel-upload"), paths)
        self.assertNotIn(("GET", "/email-campaigns/{campaign_id}/records"), paths)
        owners = [r.endpoint.__module__ for r in main.app.routes
                  if getattr(r, "path", None) == "/email-campaigns" and "GET" in (getattr(r, "methods", None) or ())]
        self.assertEqual(owners, ["routers.email_campaigns"])
        r = self.client.get("/email-campaigns", headers={"X-User-ID": str(ADMIN)})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["items"], [])
        self.assertNotIn("campaigns", r.json())  # the old route's shape
        self.assertIn(self.client.post("/email-campaigns/excel-upload", headers={"X-User-ID": str(ADMIN)}).status_code,
                      (404, 405))
        self.assertEqual(self.client.get("/email-campaigns/1/records", headers={"X-User-ID": str(ADMIN)}).status_code, 404)


if __name__ == "__main__":
    unittest.main()
