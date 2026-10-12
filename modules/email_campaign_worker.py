"""
Excel / CSV email campaigns: parse the upload, personalize each row, send from
the company's configured mailbox, and keep a per-recipient record.

Templates use {{placeholders}} — any column header (normalized: "First Name" ->
{{first_name}}) plus the aliases {{name}}, {{first_name}}, {{last_name}},
{{email}}, {{company}} and {{sender_name}}. A fallback can follow a pipe:
{{first_name|there}}. Name placeholders fall back to "there" automatically, so
"Hi {{first_name}}," never renders as "Hi ,".

Sending runs in a daemon thread per campaign. Each recipient is claimed with a
conditional UPDATE (Pending -> Sending) before it is sent, so a resumed campaign
(after a restart, or on a second app instance) never emails the same row twice.
A row stuck in Sending for 10+ minutes was interrupted mid-send; it is marked
Failed with an explanation instead of being retried blindly.
"""
import csv
import io
import os
import re
import threading
import time
from datetime import datetime, timedelta, timezone
from typing import Optional

from sqlalchemy import func, update
from sqlmodel import Session, select

import database
from database import (
    ActivityLog, Contact, EmailCampaign, EmailCampaignRecipient, FollowUp, Lead, Notification, Tenant, User,
)

MAX_FILE_BYTES = 5 * 1024 * 1024
STALE_CLAIM = timedelta(minutes=10)
_EMAIL_RE = re.compile(r"^[^@\s<>,;\"']+@[^@\s<>,;\"']+\.[A-Za-z]{2,}$")
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}|]+?)\s*(?:\|\s*([^{}]*?)\s*)?\}\}")
NAME_KEYS = {"name", "first_name", "full_name", "contact_name"}
ALIASES = ("email", "name", "first_name", "last_name", "company", "sender_name")


def max_rows() -> int:
    return int(os.getenv("BULK_EMAIL_MAX_ROWS", "2000"))


class UploadError(ValueError):
    """A user-facing problem with the uploaded file."""


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def norm_key(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")


def valid_email(value: str) -> bool:
    return bool(value) and bool(_EMAIL_RE.match(value.strip()))


# ── Parsing ──────────────────────────────────────────────────────────────────
def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if (value.hour, value.minute, value.second) == (0, 0, 0) else value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # phone numbers / ids read as 9876543210.0
    return str(value).strip()


def _table_from_rows(raw_rows) -> tuple:
    rows = [[_cell(v) for v in r] for r in raw_rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise UploadError("The file is empty.")
    header, body = rows[0], rows[1:]
    width = max(len(r) for r in rows)
    columns, seen = [], {}
    for i in range(width):
        name = header[i] if i < len(header) and header[i] else f"Column {i + 1}"
        base = name
        seen[base] = seen.get(base, 0) + 1
        if seen[base] > 1:
            name = f"{base} ({seen[base]})"
        columns.append(name)
    if len(body) > max_rows():
        raise UploadError(f"The file has {len(body)} rows; the limit is {max_rows()} per campaign. Split it into smaller files.")
    data = [{columns[i]: (r[i] if i < len(r) else "") for i in range(width)} for r in body]
    return columns, data


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise UploadError("Could not read the file's text encoding. Save it as UTF-8 CSV or .xlsx.")


def parse_upload(filename: str, data: bytes) -> dict:
    """Parse an .xlsx/.xlsm/.csv/.tsv/.txt upload into {columns, rows, detected, stats}."""
    if not data:
        raise UploadError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise UploadError("The file is larger than 5 MB.")
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm") or data[:2] == b"PK":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise UploadError("Could not open the Excel file. Make sure it is a valid .xlsx file.") from None
        try:
            sheet = next((ws for ws in wb.worksheets if ws.max_row and ws.max_row > 0), wb.active)
            columns, rows = _table_from_rows(sheet.iter_rows(values_only=True))
        finally:
            wb.close()
    elif ext == ".xls":
        raise UploadError("The old .xls format isn't supported. Open it in Excel and save as .xlsx (or .csv), then upload again.")
    elif ext in (".csv", ".tsv", ".txt", ""):
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = "\t" if ext == ".tsv" else ","
        columns, rows = _table_from_rows(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise UploadError("Unsupported file type. Upload an Excel (.xlsx) or CSV file.")
    detected = detect_mapping(columns, rows)
    return {"columns": columns, "rows": rows, "detected": detected, "stats": recipient_stats(rows, detected),
            "placeholders": placeholders_for(columns)}


def detect_mapping(columns: list, rows: list) -> dict:
    """Guess which columns hold the email, names and company."""
    keys = {c: norm_key(c) for c in columns}

    def pick(exact: tuple, contains: tuple = (), exclude: tuple = ()):
        for want in exact:
            for c, k in keys.items():
                if k == want:
                    return c
        for c, k in keys.items():
            if any(w in k for w in contains) and not any(x in k for x in exclude):
                return c
        return None

    email = pick(("email", "email_address", "e_mail", "mail", "email_id"), ("email", "e_mail"))
    if not email and rows:  # no email-looking header: the column whose values look like addresses
        best, best_hits = None, 0
        for c in columns:
            hits = sum(1 for r in rows[:50] if valid_email(r.get(c, "")))
            if hits > best_hits:
                best, best_hits = c, hits
        email = best
    return {
        "email": email,
        "first_name": pick(("first_name", "firstname", "first", "given_name", "fname"), ("first_name",)),
        "last_name": pick(("last_name", "lastname", "surname", "family_name", "lname"), ("last_name", "surname")),
        "name": pick(("name", "full_name", "contact_name", "contact", "person", "client_name", "customer_name",
                      "contact_person"), ("full_name", "contact_name", "contact_person"), ("company", "business")),
        "company": pick(("company", "company_name", "organization", "organisation", "business", "business_name",
                         "firm", "account", "account_name", "brand"), ("company", "organi", "business"), ("email",)),
    }


def placeholders_for(columns: list) -> list:
    out = list(ALIASES)
    for c in columns:
        k = norm_key(c)
        if k and k not in out:
            out.append(k)
    return out


def row_context(row: dict, mapping: dict, sender_name: str = "") -> dict:
    """Placeholder values for one row: every column (normalized) plus the aliases."""
    ctx = {norm_key(k): (v or "").strip() for k, v in row.items() if norm_key(k)}
    get = lambda field: (row.get(mapping.get(field) or "", "") or "").strip()  # noqa: E731
    first, last, full = get("first_name"), get("last_name"), get("name")
    if not full:
        full = " ".join(x for x in [first, last] if x)
    if not first and full:
        first = full.split()[0]
    if not last and full and len(full.split()) > 1:
        last = full.split()[-1]
    ctx.update({"email": get("email"), "name": full, "first_name": first, "last_name": last,
                "company": get("company"), "sender_name": sender_name or ""})
    return ctx


def unknown_placeholders(templates: list, columns: list) -> list:
    known = set(placeholders_for(columns))
    found = []
    for t in templates:
        for key, _ in _PLACEHOLDER_RE.findall(t or ""):
            k = norm_key(key)
            if k not in known and k not in found:
                found.append(k)
    return found


def render(template: str, ctx: dict) -> tuple:
    """Fill {{placeholders}}. Returns (text, [placeholders that had no value])."""
    missing = []

    def sub(m):
        key, fallback = norm_key(m.group(1)), m.group(2)
        value = ctx.get(key, "")
        if value:
            return value
        if fallback is not None and fallback != "":
            return fallback
        if key in NAME_KEYS:
            return "there"
        if key not in missing:
            missing.append(key)
        return ""

    text = _PLACEHOLDER_RE.sub(sub, template or "")
    return re.sub(r"[ \t]+([,.!?])", r"\1", text), missing


def classify_rows(rows: list, mapping: dict) -> list:
    """[(row_number, row, status, reason)] — Pending for sendable rows, Skipped with a reason otherwise."""
    out, seen = [], {}
    col = mapping.get("email")
    for i, row in enumerate(rows, start=1):
        email = (row.get(col, "") if col else "").strip()
        if not email:
            out.append((i, row, "Skipped", "No email address in this row."))
        elif not valid_email(email):
            out.append((i, row, "Skipped", f"“{email}” is not a valid email address."))
        elif email.lower() in seen:
            out.append((i, row, "Skipped", f"Duplicate of row {seen[email.lower()]}."))
        else:
            seen[email.lower()] = i
            out.append((i, row, "Pending", None))
    return out


def recipient_stats(rows: list, mapping: dict) -> dict:
    statuses = classify_rows(rows, mapping)
    return {
        "rows": len(rows),
        "sendable": sum(1 for s in statuses if s[2] == "Pending"),
        "missing_email": sum(1 for s in statuses if s[3] and s[3].startswith("No email")),
        "invalid_email": sum(1 for s in statuses if s[3] and "not a valid" in s[3]),
        "duplicates": sum(1 for s in statuses if s[3] and s[3].startswith("Duplicate")),
    }


# ── AI personalization ───────────────────────────────────────────────────────
def personalize(subject: str, body: str, row: dict, instructions: str = "", sender_name: str = "",
                company: str = "") -> dict:
    """Let the AI tailor an already-merged email to this row's details. Raises EmailAIError on failure."""
    from modules import email_ai
    system = (
        f"You personalize one outreach email for one recipient on behalf of {sender_name or 'a salesperson'}"
        f"{f' at {company}' if company else ''}. Keep the email's message, offer, call to action, sign-off and "
        "approximate length. Adapt the greeting and one or two sentences so they reference the recipient's own "
        "details from their spreadsheet row (role, company, city, industry, notes...). Use only facts present in the "
        "row or the email — never invent results, prices, dates or names. Plain text, no markdown, no bracketed "
        f"placeholders.\n{email_ai.UNTRUSTED}\n"
        'Respond as JSON: {"subject": "...", "body": "..."}'
    )
    row_text = "\n".join(f"{k}: {email_ai._clip(v, 300)}" for k, v in row.items() if str(v or "").strip())
    user = (f"<crm_context>\nRecipient's spreadsheet row:\n{row_text}\n</crm_context>\n\n"
            f"Email to personalize:\nSubject: {subject}\n\n{body}")
    if instructions:
        user += f"\n\nSalesperson's instructions (follow these): {email_ai._clip(instructions, 1000)}"
    data = email_ai._chat_json(system, user, model=email_ai.TRIAGE_MODEL, temperature=0.5, max_tokens=900, timeout=45)
    new_body = str(data.get("body") or "").strip()
    if not new_body:
        raise email_ai.EmailAIError("The AI returned an empty email.")
    return {"subject": (str(data.get("subject") or "").strip() or subject)[:300], "body": new_body[:8000]}


def compose(campaign: EmailCampaign, row: dict, sender_name: str, company: str, ai=True) -> tuple:
    """(subject, body, note) for one row: template merge, then optional AI personalization."""
    ctx = row_context(row, campaign.mapping or {}, sender_name)
    subject, _ = render(campaign.subject_template, ctx)
    body, _ = render(campaign.body_template, ctx)
    note = None
    if ai and campaign.ai_personalize:
        from modules.email_ai import EmailAIError
        try:
            p = personalize(subject, body, row, campaign.ai_instructions or "", sender_name, company)
            subject, body = p["subject"], p["body"]
        except EmailAIError as e:
            note = f"AI personalization unavailable ({e}); sent the merged template."
    return subject.strip()[:500], body, note


# ── Sending ──────────────────────────────────────────────────────────────────
_running = set()
_running_lock = threading.Lock()


def refresh_counts(session: Session, campaign: EmailCampaign):
    counts = dict(session.exec(
        select(EmailCampaignRecipient.status, func.count()).where(EmailCampaignRecipient.campaign_id == campaign.id)
        .group_by(EmailCampaignRecipient.status).execution_options(skip_tenant=True)
    ).all())
    campaign.total = sum(counts.values())
    campaign.sent_count = counts.get("Sent", 0)
    campaign.failed_count = counts.get("Failed", 0)
    campaign.skipped_count = counts.get("Skipped", 0)
    session.add(campaign)
    return counts


def _link_or_create_lead(session: Session, campaign: EmailCampaign, r: EmailCampaignRecipient) -> tuple:
    """(lead_id, client_id, created) — an existing lead/client by email, else a new lead if enabled."""
    from modules.inbox_sync import match_address
    links = match_address(session, r.email, campaign.tenant_id)
    lead_id, client_id = links.get("lead_id"), links.get("client_id")
    created = False
    if not lead_id and not client_id and campaign.create_leads:
        domain = r.email.split("@")[-1]
        lead = Lead(tenant_id=campaign.tenant_id, company_name=(r.company or r.name or domain)[:255], email=r.email,
                    source="Email Campaign", status="Contacted", owner_id=campaign.created_by,
                    notes=f"Imported from email campaign “{campaign.name}”.")
        session.add(lead)
        session.flush()
        lead_id, created = lead.id, True
        if r.name:
            parts = r.name.split()
            session.add(Contact(tenant_id=campaign.tenant_id, first_name=parts[0][:100],
                                last_name=" ".join(parts[1:])[:100] or None, full_name=r.name[:255], email=r.email,
                                lead_id=lead.id, owner_id=campaign.created_by))
    return lead_id, client_id, created


def _discard_new_lead(session: Session, lead_id: Optional[int]):
    """A send failed: don't leave behind a lead created only for that recipient."""
    lead = session.get(Lead, lead_id) if lead_id else None
    if not lead:
        return
    for c in session.exec(select(Contact).where(Contact.lead_id == lead_id)).all():
        session.delete(c)
    session.flush()
    session.delete(lead)


def _after_send(session: Session, campaign: EmailCampaign, r: EmailCampaignRecipient, subject: str):
    from modules.followups_core import at_local_hour
    if r.lead_id or r.client_id:
        session.add(ActivityLog(
            tenant_id=campaign.tenant_id, userId=campaign.created_by, clientId=r.client_id, lead_id=r.lead_id,
            action=f"Campaign email sent: {subject}"[:255], method="Email",
            content=f"Sent via campaign “{campaign.name}” to {r.email}", details=(r.body or "")[:500],
        ))
    if r.lead_id:
        lead = session.get(Lead, r.lead_id)
        if lead:
            lead.last_activity = f"Campaign email sent: {subject}"[:500]
            if (lead.status or "New") in ("New", "Generated"):
                lead.status = "Contacted"
            session.add(lead)
    if campaign.followup_days:
        due = at_local_hour((utcnow() + timedelta(days=int(campaign.followup_days))).date())
        # The record's own salesperson follows up; the campaign sender only for new / unowned contacts.
        from modules.followups_core import record_info
        owner = record_info(session, r.lead_id, r.client_id).get("owner_id") if (r.lead_id or r.client_id) else None
        fu = FollowUp(
            tenant_id=campaign.tenant_id, lead_id=r.lead_id, client_id=r.client_id, contact_name=r.name,
            contact_email=r.email, company_name=r.company, title=f"Follow up on “{subject}”"[:500],
            notes=(f"Sent “{subject}” on {utcnow().date().isoformat()} via the email campaign “{campaign.name}”. "
                   "If they haven't replied, send a short personalized follow-up."),
            channel="Email", due_at=due, source="campaign", source_ref=f"campaign-{campaign.id}",
            owner_id=owner or campaign.created_by, created_by=campaign.created_by,
        )
        session.add(fu)
        session.flush()
        r.follow_up_id = fu.id


def _claim(session: Session, recipient_id: int) -> bool:
    res = session.execute(
        update(EmailCampaignRecipient)
        .where(EmailCampaignRecipient.id == recipient_id, EmailCampaignRecipient.status == "Pending")
        .values(status="Sending", claimed_at=utcnow(), attempts=EmailCampaignRecipient.attempts + 1)
    )
    session.commit()
    return res.rowcount == 1


def _finish(session: Session, campaign: EmailCampaign, status: str, error: Optional[str] = None):
    campaign.status = status
    campaign.completed_at = utcnow()
    if error:
        campaign.last_error = error
    counts = refresh_counts(session, campaign)
    if status == "Completed":
        message = f"{counts.get('Sent', 0)} sent, {counts.get('Failed', 0)} failed, {counts.get('Skipped', 0)} skipped."
    else:
        message = error or status
    if campaign.created_by:
        session.add(Notification(tenant_id=campaign.tenant_id, user_id=campaign.created_by,
                                 title=f"Email campaign “{campaign.name}” {status.lower()}"[:255], message=message,
                                 type="info", link=f"/email-agent?campaign={campaign.id}"))
    session.commit()


def run_campaign(campaign_id: int, engine=None, sender=None, delay: Optional[float] = None, ai=True):
    """Send every Pending recipient of a campaign, one at a time. Synchronous (the thread target)."""
    from modules import company_mailer
    engine = engine or database.engine
    delay = float(os.getenv("BULK_EMAIL_DELAY_SECONDS", "3")) if delay is None else delay
    with Session(engine) as s:
        c = s.get(EmailCampaign, campaign_id)
        if not c or c.status not in ("Queued", "Sending"):
            return
        if c.status == "Queued":
            c.status, c.started_at = "Sending", utcnow()
            s.add(c)
            s.commit()
        try:
            account = company_mailer.resolve_account(s, c.created_by, c.tenant_id)
        except company_mailer.MailAccountError as e:
            _finish(s, c, "Failed", str(e))
            return
        signature = company_mailer.user_signature(s, c.created_by)
        creator = s.get(User, c.created_by) if c.created_by else None
        tenant = s.get(Tenant, c.tenant_id) if c.tenant_id else None
        sender_name = account.from_name or (creator.name if creator else "") or ""
        company = (tenant.business_name or tenant.name) if tenant else ""

    while True:
        with Session(engine) as s:
            c = s.get(EmailCampaign, campaign_id)
            if not c or c.status != "Sending":
                return  # cancelled
            r = s.exec(select(EmailCampaignRecipient).where(
                EmailCampaignRecipient.campaign_id == campaign_id, EmailCampaignRecipient.status == "Pending",
            ).order_by(EmailCampaignRecipient.id).limit(1)).first()
            if r is None:
                if not s.exec(select(EmailCampaignRecipient.id).where(
                        EmailCampaignRecipient.campaign_id == campaign_id, EmailCampaignRecipient.status == "Sending")).first():
                    _finish(s, c, "Completed")
                return
            if not _claim(s, r.id):
                continue
            s.refresh(r)
            subject, body, note = compose(c, r.data or {}, sender_name, company, ai=ai)
            body = company_mailer.apply_signature(body, signature)
            r.subject, r.body = subject, body
            new_lead = False
            try:
                r.lead_id, r.client_id, new_lead = _link_or_create_lead(s, c, r)
                s.add(r)
                s.commit()
                sent = company_mailer.send_and_record(
                    s, account, tenant_id=c.tenant_id, to=r.email, subject=subject, body=body,
                    lead_id=r.lead_id, client_id=r.client_id,
                    meta={"source": "email_campaign", "campaign_id": c.id, "recipient_id": r.id,
                          "sent_by_user_id": c.created_by},
                    **({"sender": sender} if sender else {}),
                )
            except Exception as e:
                import smtplib
                s.rollback()
                r = s.get(EmailCampaignRecipient, r.id)
                if new_lead:
                    _discard_new_lead(s, r.lead_id)
                    r.lead_id = None
                r.status, r.error = "Failed", company_mailer.friendly_smtp_error(e)
                s.add(r)
                s.commit()
                if isinstance(e, smtplib.SMTPAuthenticationError):
                    # Every remaining row would fail the same way — stop and tell the user.
                    _finish(s, s.get(EmailCampaign, campaign_id), "Failed", r.error)
                    return
                refresh_counts(s, s.get(EmailCampaign, campaign_id))
                s.commit()
            else:
                s.refresh(r)
                r.status, r.sent_at, r.sent_email_id, r.error = "Sent", utcnow(), sent.id, note
                _after_send(s, c, r, subject)
                s.add(r)
                refresh_counts(s, c)
                s.commit()
        if delay:
            time.sleep(delay)


def _run_safe(campaign_id: int, engine=None):
    try:
        run_campaign(campaign_id, engine)
    except Exception as e:
        print(f"[Campaigns] campaign {campaign_id} crashed: {e}")
        try:
            with Session(engine or database.engine) as s:
                c = s.get(EmailCampaign, campaign_id)
                if c and c.status == "Sending":
                    c.last_error = f"Sending stopped unexpectedly: {e}"[:500]
                    s.add(c)
                    s.commit()
        except Exception:
            pass
    finally:
        with _running_lock:
            _running.discard(campaign_id)


def start_campaign(campaign_id: int, engine=None) -> bool:
    """Run a campaign in the background (no-op if this process is already sending it)."""
    with _running_lock:
        if campaign_id in _running:
            return False
        _running.add(campaign_id)
    threading.Thread(target=_run_safe, args=(campaign_id, engine), daemon=True, name=f"campaign-{campaign_id}").start()
    return True


def resume_campaigns(engine=None) -> int:
    """Called by the follow-up scheduler: fail rows interrupted mid-send, restart unfinished campaigns."""
    engine = engine or database.engine
    resumed = 0
    with Session(engine) as s:
        stale = s.exec(select(EmailCampaignRecipient).where(
            EmailCampaignRecipient.status == "Sending", EmailCampaignRecipient.claimed_at < utcnow() - STALE_CLAIM,
        ).execution_options(skip_tenant=True).limit(200)).all()
        for r in stale:
            r.status = "Failed"
            r.error = "Interrupted while sending (server restart). It may not have been delivered — use Retry failed to resend."
            s.add(r)
        s.commit()
        ids = s.exec(select(EmailCampaign.id).where(EmailCampaign.status.in_(("Queued", "Sending")))
                     .execution_options(skip_tenant=True)).all()
    for cid in ids:
        with _running_lock:
            busy = cid in _running
        if not busy and start_campaign(cid, engine):
            resumed += 1
    return resumed
